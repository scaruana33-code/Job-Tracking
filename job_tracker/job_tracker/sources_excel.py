"""
Connector — read from AND write to Executive_Job_Search_Tracker.xlsx.

Maps YOUR tracker columns to the canonical schema, so the spreadsheet is both:
  * an INPUT source (rows you log by hand join the DB/Gmail/LinkedIn feeds), and
  * a live OUTPUT (the unified fact table is written back, and the Dashboard +
    Source Analysis tabs are recomputed on every refresh).

Tracker "Applications" columns:
    Date Applied | Week | Company | Job Title | Location |
    Source (ATS/LinkedIn/Networking/Recruiter) | Hiring Manager |
    Application Status | Interview Count | Current Stage |
    1st Round Date | 2nd Round Date | 3rd Round Date | Final Round Date |
    Offer Received (Y/N) | Offer Date | Notes
"""
from __future__ import annotations

import hashlib
import logging
import os
from typing import Optional

import pandas as pd

from .schema import CANONICAL_COLUMNS, STAGE_ORDER

log = logging.getLogger(__name__)

APPLICATIONS_SHEET = "Applications"
DASHBOARD_SHEET = "Dashboard"
SOURCE_SHEET = "Source Analysis"

_SOURCE_MAP = {
    "ats": "ats", "linkedin": "linkedin", "networking": "networking",
    "referral": "networking", "recruiter": "recruiter",
}
_STAGE_TO_DATECOL = {
    "phone_1": "1st Round Date",
    "interview_2": "2nd Round Date",
    "interview_3": "3rd Round Date",
    "onsite_final": "Final Round Date",
}
_INTERVIEW_STAGES = ["phone_1", "interview_2", "interview_3", "onsite_final", "offer"]


def _hash_id(*parts) -> str:
    return "xlsx:" + hashlib.sha1("|".join(str(p or "") for p in parts).encode()).hexdigest()[:16]


def _norm_source(label) -> str:
    if not isinstance(label, str):
        return "ats"
    return _SOURCE_MAP.get(label.strip().lower(), "ats")


def _has_value(v) -> bool:
    """True only for a real, non-empty cell (empty Excel cells read as NaN)."""
    if v is None:
        return False
    if isinstance(v, float) and pd.isna(v):
        return False
    return str(v).strip() != ""


def _infer_stage(row) -> str:
    offer_flag = str(row.get("Offer Received (Y/N)", "")).strip().lower() in ("y", "yes", "true", "1")
    if offer_flag or _has_value(row.get("Offer Date")):
        return "offer"
    for stage in ("onsite_final", "interview_3", "interview_2", "phone_1"):
        if _has_value(row.get(_STAGE_TO_DATECOL[stage])):
            return stage
    cur = str(row.get("Current Stage", "")).strip().lower()
    if cur and cur != "nan":
        # match the longest stage label first to avoid "applied" swallowing others
        for key in sorted(STAGE_ORDER, key=lambda k: -len(k)):
            if key.replace("_", " ") in cur or key in cur:
                return key
    return "applied"


# ---- READ --------------------------------------------------------------- #
def pull(xlsx_path: str) -> pd.DataFrame:
    if not os.path.exists(xlsx_path):
        log.info("Tracker not found at %s (skipping Excel source).", xlsx_path)
        return pd.DataFrame(columns=CANONICAL_COLUMNS)

    raw = pd.read_excel(xlsx_path, sheet_name=APPLICATIONS_SHEET)
    raw = raw[raw["Company"].notna() & (raw["Company"].astype(str).str.strip() != "")]
    if raw.empty:
        log.info("Tracker has no real application rows yet.")
        return pd.DataFrame(columns=CANONICAL_COLUMNS)

    rows = []
    for _, r in raw.iterrows():
        stage = _infer_stage(r)
        applied = pd.to_datetime(r.get("Date Applied"), errors="coerce", utc=True)
        rows.append(
            {
                "application_id": _hash_id(r.get("Company"), r.get("Job Title"), applied),
                "company": r.get("Company"),
                "role": r.get("Job Title"),
                "applied_at": applied,
                "source": _norm_source(r.get("Source (ATS/LinkedIn/Networking/Recruiter)")),
                "stage": stage,
                "stage_level": STAGE_ORDER.get(stage, 0),
                "offer": stage == "offer",
                "origin": "excel",
                "notes": r.get("Notes"),
            }
        )
    log.info("Excel tracker -> %d applications", len(rows))
    return pd.DataFrame(rows)[CANONICAL_COLUMNS]


# ---- WRITE -------------------------------------------------------------- #
def write_back(fact: pd.DataFrame, xlsx_path: str, out_path: Optional[str] = None) -> str:
    import openpyxl

    out_path = out_path or xlsx_path
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = APPLICATIONS_SHEET
        # seed the header row so writes have somewhere to go
        wb[APPLICATIONS_SHEET].append(
            [
                "Date Applied", "Week", "Company", "Job Title", "Location",
                "Source (ATS/LinkedIn/Networking/Recruiter)", "Hiring Manager",
                "Application Status", "Interview Count", "Current Stage",
                "1st Round Date", "2nd Round Date", "3rd Round Date",
                "Final Round Date", "Offer Received (Y/N)", "Offer Date", "Notes",
            ]
        )
        wb.create_sheet(DASHBOARD_SHEET)
        wb.create_sheet(SOURCE_SHEET)

    _write_applications(wb, fact)
    _write_dashboard(wb, fact)
    _write_source_analysis(wb, fact)

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    log.info("Wrote refreshed tracker to %s", out_path)
    return out_path


def _fmt_date(ts):
    if ts is None or pd.isna(ts):
        return None
    return pd.Timestamp(ts).strftime("%Y-%m-%d")


def _write_applications(wb, fact: pd.DataFrame) -> None:
    ws = wb[APPLICATIONS_SHEET]
    headers = [c.value for c in ws[1]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    idx = {h: i for i, h in enumerate(headers)}

    for _, r in fact.sort_values("applied_at").iterrows():
        line = [None] * len(headers)

        def put(col, val):
            if col in idx:
                line[idx[col]] = val

        put("Date Applied", _fmt_date(r["applied_at"]))
        put("Week", _fmt_date(r.get("week")))
        put("Company", r["company"])
        put("Job Title", r["role"])
        put("Source (ATS/LinkedIn/Networking/Recruiter)", str(r["source"]).title())
        put("Application Status", "Offer" if r["offer"] else str(r["stage"]).replace("_", " ").title())
        put("Interview Count", 1 if r["stage"] in _INTERVIEW_STAGES else 0)
        put("Current Stage", str(r["stage"]).replace("_", " ").title())
        lvl = STAGE_ORDER.get(r["stage"], 0)
        for stage, colname in _STAGE_TO_DATECOL.items():
            put(colname, "X" if lvl >= STAGE_ORDER[stage] else None)
        put("Offer Received (Y/N)", "Y" if r["offer"] else "N")
        put("Notes", r.get("notes"))
        ws.append(line)


def _write_dashboard(wb, fact: pd.DataFrame) -> None:
    ws = wb[DASHBOARD_SHEET] if DASHBOARD_SHEET in wb.sheetnames else wb.create_sheet(DASHBOARD_SHEET)
    total = len(fact)
    interviews = int(fact["stage"].isin(_INTERVIEW_STAGES).sum()) if total else 0
    offers = int(fact["offer"].sum()) if total else 0
    by_src = fact["source"].value_counts().to_dict() if total else {}
    metrics = [
        ("Metric", "Value"),
        ("Total Applications", total),
        ("Total Interviews", interviews),
        ("Offers Received", offers),
        ("ATS Applications", by_src.get("ats", 0)),
        ("LinkedIn Applications", by_src.get("linkedin", 0)),
        ("Networking Applications", by_src.get("networking", 0)),
        ("Recruiter Initiated", by_src.get("recruiter", 0)),
        ("Interview Conversion Rate", round(interviews / total, 3) if total else 0),
        ("Offer Conversion Rate", round(offers / total, 3) if total else 0),
    ]
    _overwrite(ws, metrics)


def _write_source_analysis(wb, fact: pd.DataFrame) -> None:
    ws = wb[SOURCE_SHEET] if SOURCE_SHEET in wb.sheetnames else wb.create_sheet(SOURCE_SHEET)
    data = [("Source", "Applications", "Interviews", "Offers")]
    for src in ["ATS", "LinkedIn", "Networking", "Recruiter"]:
        g = fact[fact["source"] == src.lower()] if not fact.empty else fact
        apps = len(g)
        iv = int(g["stage"].isin(_INTERVIEW_STAGES).sum()) if apps else 0
        off = int(g["offer"].sum()) if apps else 0
        data.append((src, apps, iv, off))
    _overwrite(ws, data)


def _overwrite(ws, rows) -> None:
    if ws.max_row >= 1:
        ws.delete_rows(1, ws.max_row)
    for row in rows:
        ws.append(list(row))
